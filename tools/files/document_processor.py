"""DOCX attachment processing."""

from __future__ import annotations

import io

from docx import Document

from tools.files.models import AttachmentMetadata, ProcessedFileContext, ProcessingStatus


def process_docx(metadata: AttachmentMetadata, content: bytes) -> ProcessedFileContext:
    """Extract paragraphs, headings, and tables from a DOCX file."""
    try:
        document = Document(io.BytesIO(content))
    except Exception:
        return _failed(metadata, "The DOCX document could not be read.")

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    headings = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.style and "Heading" in paragraph.style.name]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"rows": rows})

    if not paragraphs and not tables:
        return _failed(metadata, "The DOCX document contains no readable content.")

    return ProcessedFileContext(
        attachment_id=metadata.attachment_id,
        filename=metadata.original_filename,
        file_type="docx",
        processing_status=ProcessingStatus.READY,
        summary=f"DOCX document with {len(paragraphs)} paragraph(s) and {len(tables)} table(s).",
        paragraphs=paragraphs,
        tables=tables,
        file_schema={"headings": headings},
        full_data={"text": "\n".join(paragraphs)},
        source_references=[metadata.original_filename],
    )


def _failed(metadata: AttachmentMetadata, message: str) -> ProcessedFileContext:
    return ProcessedFileContext(
        attachment_id=metadata.attachment_id,
        filename=metadata.original_filename,
        file_type="docx",
        processing_status=ProcessingStatus.FAILED,
        error_message=message,
        source_references=[metadata.original_filename],
    )
