"""Robust Excel reading shared by ingestion, analysis, and UI preview."""

from __future__ import annotations

import io
from typing import Any

import pandas as pd
from openpyxl import load_workbook


def read_excel_sheets(content: bytes) -> dict[str, list[dict[str, Any]]]:
    """
    Read every sheet into list-of-dict rows.

    Logistics Excel exports often declare incorrect worksheet dimensions
    (for example max_row=2 while thousands of rows exist). openpyxl's
    read_only mode trusts those dimensions and returns only the first row.

    Strategy:
    1. openpyxl in normal (non-read-only) mode — uses real cell data
    2. pandas fallback
    """
    try:
        return _read_with_openpyxl(content)
    except Exception:
        return _read_with_pandas(content)


def read_excel_sheet(
    content: bytes,
    *,
    sheet_name: str | None = None,
    max_rows: int | None = None,
) -> tuple[list[str], list[dict[str, Any]], list[str], str | None]:
    """Return columns, rows, sheet names, and the active sheet name."""
    sheets = read_excel_sheets(content)
    sheet_names = list(sheets.keys())
    if not sheet_names:
        return [], [], [], None

    active = sheet_name if sheet_name in sheets else sheet_names[0]
    rows = sheets[active]
    preview_rows = rows if max_rows is None else rows[:max_rows]
    columns = list(preview_rows[0].keys()) if preview_rows else list(rows[0].keys()) if rows else []
    return columns, preview_rows, sheet_names, active


def _read_with_openpyxl(content: bytes) -> dict[str, list[dict[str, Any]]]:
    # IMPORTANT: do not use read_only=True — it truncates these exports.
    workbook = load_workbook(filename=io.BytesIO(content), data_only=True, read_only=False)
    sheets: dict[str, list[dict[str, Any]]] = {}
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            raw_rows = [
                tuple(row)
                for row in worksheet.iter_rows(values_only=True)
            ]
            sheets[str(sheet_name)] = _rows_to_records(raw_rows)
    finally:
        workbook.close()
    return sheets


def _read_with_pandas(content: bytes) -> dict[str, list[dict[str, Any]]]:
    buffer = io.BytesIO(content)
    try:
        frames = pd.read_excel(buffer, sheet_name=None, dtype=object, engine="openpyxl")
    except Exception:
        buffer.seek(0)
        frames = pd.read_excel(buffer, sheet_name=None, dtype=object)

    sheets: dict[str, list[dict[str, Any]]] = {}
    for sheet_name, frame in (frames or {}).items():
        sheets[str(sheet_name)] = _frame_to_records(frame)
    return sheets


def _rows_to_records(raw_rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    if not raw_rows:
        return []

    # Skip leading fully empty rows before the header.
    start = 0
    while start < len(raw_rows) and all(cell is None or str(cell).strip() == "" for cell in raw_rows[start]):
        start += 1
    if start >= len(raw_rows):
        return []

    header_cells = raw_rows[start]
    header: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(header_cells):
        name = str(cell).strip() if cell is not None and str(cell).strip() else f"column_{index + 1}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        header.append(name if count == 0 else f"{name}_{count + 1}")

    records: list[dict[str, Any]] = []
    for row in raw_rows[start + 1 :]:
        if row is None:
            continue
        record: dict[str, Any] = {}
        for index, column in enumerate(header):
            value = row[index] if index < len(row) else None
            record[column] = _normalize_cell(value)
        if any(value is not None and str(value).strip() != "" for value in record.values()):
            records.append(record)
    return records


def _frame_to_records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    if frame is None or frame.empty:
        return []

    cleaned = frame.dropna(axis=0, how="all").dropna(axis=1, how="all")
    if cleaned.empty:
        return []

    cleaned.columns = [
        str(column).strip()
        if column is not None and str(column).strip() and not str(column).startswith("Unnamed")
        else f"column_{index + 1}"
        for index, column in enumerate(cleaned.columns)
    ]
    seen: dict[str, int] = {}
    unique_columns: list[str] = []
    for column in cleaned.columns:
        count = seen.get(column, 0)
        seen[column] = count + 1
        unique_columns.append(column if count == 0 else f"{column}_{count + 1}")
    cleaned.columns = unique_columns

    records: list[dict[str, Any]] = []
    for row in cleaned.to_dict(orient="records"):
        normalized = {key: _normalize_cell(value) for key, value in row.items()}
        if any(value is not None and str(value).strip() != "" for value in normalized.values()):
            records.append(normalized)
    return records


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    try:
        if isinstance(value, float) and pd.isna(value):
            return None
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ", timespec="seconds")
        except TypeError:
            return value.isoformat()
    return value
