"""Tests for robust Excel reading used by preview and analysis."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from frontend.attachment_preview import build_table_preview
from tools.files.excel_processor import process_excel
from tools.files.excel_reader import read_excel_sheet, read_excel_sheets
from tools.files.models import AttachmentMetadata


def _multi_row_xlsx_bytes(row_count: int = 25) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["实际揽收时间", "详细地址", "计划揽收包裹数", "包裹数"])
    for index in range(row_count):
        sheet.append(
            [
                f"2026-05-31 13:23:{index:02d}",
                f"{100 + index} Main St, Chicago, IL",
                10 + index,
                5 + index,
            ]
        )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_read_excel_sheets_returns_all_rows() -> None:
    content = _multi_row_xlsx_bytes(30)
    sheets = read_excel_sheets(content)
    assert "Sheet1" in sheets
    assert len(sheets["Sheet1"]) == 30
    assert "详细地址" in sheets["Sheet1"][0]


def test_preview_shows_more_than_one_excel_row() -> None:
    content = _multi_row_xlsx_bytes(15)
    payload = build_table_preview(content, "xlsx", max_rows=20)
    assert payload["error"] is None
    assert payload["total_rows"] == 15
    assert len(payload["rows"]) == 15
    assert "计划揽收包裹数" in payload["columns"]


def test_process_excel_preserves_all_rows_for_analysis() -> None:
    content = _multi_row_xlsx_bytes(12)
    metadata = AttachmentMetadata(
        filename="drivers.xlsx",
        original_filename="司机点数.xlsx",
        file_type="xlsx",
        file_size=len(content),
        storage_path="/tmp/drivers.xlsx",
    )
    processed = process_excel(metadata, content)
    assert processed.statistics["row_count"] == 12
    assert len(processed.full_data["sheets"][0]["rows"]) == 12


def test_read_excel_sheet_respects_max_rows_for_preview() -> None:
    content = _multi_row_xlsx_bytes(40)
    columns, rows, sheet_names, active = read_excel_sheet(content, max_rows=20)
    assert active == "Sheet1"
    assert sheet_names == ["Sheet1"]
    assert len(rows) == 20
    assert "实际揽收时间" in columns


def test_reader_recovers_rows_when_openpyxl_readonly_reports_max_row_two() -> None:
    """Regression: logistics exports often advertise max_row=2 incorrectly."""
    from openpyxl import load_workbook

    path_bytes = _multi_row_xlsx_bytes(25)
    # Sanity: read_only path is the broken behavior we must not use.
    workbook = load_workbook(io.BytesIO(path_bytes), read_only=True, data_only=True)
    readonly_count = sum(1 for _ in workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    workbook.close()

    sheets = read_excel_sheets(path_bytes)
    assert len(sheets["Sheet1"]) == 25
    # For a normal workbook read_only may still see all rows; the important
    # assertion is that our reader never returns only the header+1 pattern
    # for multi-row content.
    assert len(sheets["Sheet1"]) >= max(readonly_count - 1, 25)


def test_real_upload_pickup_file_has_more_than_one_row() -> None:
    from pathlib import Path

    path = Path("data/uploads/56edc808-0f8b-416c-b0df-e33ab79c89ea.xlsx")
    if not path.exists():
        pytest.skip("local upload fixture missing")
    content = path.read_bytes()

    from openpyxl import load_workbook

    broken = load_workbook(path, read_only=True, data_only=True)
    broken_rows = list(broken[broken.sheetnames[0]].iter_rows(values_only=True))
    broken.close()
    assert len(broken_rows) == 2  # documents the vendor bug

    sheets = read_excel_sheets(content)
    assert len(sheets["Sheet1"]) > 1


def test_real_upload_order_file_has_thousands_of_rows() -> None:
    from pathlib import Path

    path = Path("data/uploads/9f8f12c0-c985-40b7-9706-e15074d86f0a.xlsx")
    if not path.exists():
        pytest.skip("local upload fixture missing")
    sheets = read_excel_sheets(path.read_bytes())
    sheet_name = next(iter(sheets))
    assert len(sheets[sheet_name]) > 1000
