"""Tests for attachment upload and file intelligence pipeline."""

from __future__ import annotations

import io
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from api.server import app
from core.error_handler import handle_error
from core.errors import FileError
from core.agent import GOFOAgent
from core.models import QueryResponse
from tools.files.attachment_memory import AttachmentMemory
from tools.files.csv_processor import process_csv
from tools.files.document_processor import process_docx
from tools.files.excel_processor import process_excel
from tools.files.image_processor import process_image
from tools.files.models import AttachmentMetadata, ProcessingStatus
from tools.files.pdf_processor import process_pdf
from tools.files.service import AttachmentService
from tools.files.source_router import DataSource, plan_data_sources
from tools.files.storage import AttachmentStore
from tools.files.text_processor import process_text
from tools.files.validator import validate_upload
from tools.context.file_context_builder import build_file_context, resolve_file_reference
from tools.files import analyzer as file_analyzer


CSV_CONTENT = """hub,pickup_count,package_count,delayed_pickups
Chicago Hub,10,120,7
New York Hub,100,900,2
Atlanta Hub,25,210,4
""".encode()

MONDAY_CSV = """hub,pickup_count,package_count
Chicago Hub,10,120
New York Hub,100,900
""".encode()

TUESDAY_CSV = """hub,pickup_count,package_count
Chicago Hub,15,150
New York Hub,95,850
""".encode()


def _metadata(filename: str, file_type: str, size: int = 1) -> AttachmentMetadata:
    return AttachmentMetadata(
        filename=filename,
        original_filename=filename,
        file_type=file_type,
        file_size=size,
        storage_path=f"/tmp/{filename}",
    )


def _make_excel_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["driver_name", "hub", "package_count"])
    sheet.append(["Drew Nguyen", "ORD Hub", 56])
    sheet.append(["Alex Chen", "Chicago Hub", 12])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Failed Pickup SOP", level=1)
    document.add_paragraph("Failed pickups must be reported within 2 hours.")
    document.add_paragraph("Drivers should contact dispatch immediately.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _make_png_bytes() -> bytes:
    image = Image.new("RGB", (120, 80), color=(30, 144, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def temp_store(tmp_path: Path) -> AttachmentStore:
    return AttachmentStore(base_dir=tmp_path / "uploads")


def test_validate_upload_rejects_empty_file() -> None:
    with pytest.raises(Exception, match="empty"):
        validate_upload(filename="report.csv", content_type="text/csv", file_size=0, content=b"")


def test_validate_upload_rejects_unsupported_extension() -> None:
    with pytest.raises(Exception, match="not supported"):
        validate_upload(filename="virus.exe", content_type="application/octet-stream", file_size=4, content=b"MZ\x00")


def test_validate_upload_rejects_oversized_file(monkeypatch) -> None:
    monkeypatch.setattr("config.MAX_UPLOAD_SIZE_BYTES", 10)
    with pytest.raises(Exception, match="size limit"):
        validate_upload(
            filename="large.csv",
            content_type="text/csv",
            file_size=20,
            content=b"x" * 20,
        )


def test_csv_processor_extracts_schema_and_full_rows() -> None:
    metadata = AttachmentMetadata(
        filename="safe.csv",
        original_filename="pickup_report.csv",
        file_type="csv",
        file_size=len(CSV_CONTENT),
        storage_path="/tmp/pickup_report.csv",
    )
    processed = process_csv(metadata, CSV_CONTENT)
    assert processed.processing_status == ProcessingStatus.READY
    assert processed.statistics["row_count"] == 3
    assert "hub" in processed.file_schema["columns"]
    assert len(processed.full_data["rows"]) == 3


def test_excel_processor_extracts_sheets_and_rows() -> None:
    content = _make_excel_bytes()
    processed = process_excel(_metadata("pickup_report.xlsx", "xlsx", len(content)), content)
    assert processed.processing_status == ProcessingStatus.READY
    assert processed.file_schema["active_sheet"] == "Sheet1"
    assert processed.statistics["row_count"] == 2
    assert "Drew Nguyen" in str(processed.sample_rows)


def test_pdf_processor_extracts_pages() -> None:
    content = _make_pdf_bytes()
    processed = process_pdf(_metadata("operations_report.pdf", "pdf", len(content)), content)
    assert processed.processing_status == ProcessingStatus.READY
    assert len(processed.pages) == 1


def test_docx_processor_extracts_paragraphs() -> None:
    content = _make_docx_bytes()
    processed = process_docx(_metadata("sop.docx", "docx", len(content)), content)
    assert processed.processing_status == ProcessingStatus.READY
    assert any("Failed pickups" in paragraph for paragraph in processed.paragraphs)


def test_text_processor_extracts_sections() -> None:
    content = b"Section One\n\nFailed pickups must be escalated."
    processed = process_text(_metadata("notes.txt", "txt", len(content)), content)
    assert processed.processing_status == ProcessingStatus.READY
    assert len(processed.paragraphs) >= 1


@patch("tools.files.image_processor._analyze_image")
def test_image_processor_uses_multimodal_analysis(mock_analyze: MagicMock) -> None:
    mock_analyze.return_value = {
        "description": "Dashboard shows delayed pickups rising at Chicago Hub.",
        "extracted_text": "Chicago Hub 82%",
        "observed_metrics": ["82% completion"],
        "analysis": "Operational delay issue visible.",
    }
    content = _make_png_bytes()
    processed = process_image(_metadata("dashboard.png", "png", len(content)), content)
    assert processed.processing_status == ProcessingStatus.READY
    assert "Chicago Hub" in processed.image_analysis["description"]
    mock_analyze.assert_called_once()


def test_corrupted_excel_returns_failed_status() -> None:
    processed = process_excel(_metadata("broken.xlsx", "xlsx"), b"not-a-real-workbook")
    assert processed.processing_status == ProcessingStatus.FAILED
    assert "could not be read" in (processed.error_message or "").lower()


def test_attachment_service_upload_and_process(temp_store: AttachmentStore) -> None:
    service = AttachmentService(store=temp_store)
    metadata = service.upload(
        filename="pickup_report.csv",
        content=CSV_CONTENT,
        content_type="text/csv",
        conversation_id="session-1",
    )
    processed = service.process(metadata.attachment_id)
    assert processed.file_type == "csv"
    assert processed.statistics["row_count"] == 3


def test_upload_persistence_survives_store_reload(temp_store: AttachmentStore) -> None:
    service = AttachmentService(store=temp_store)
    metadata = service.upload(
        filename="pickup_report.csv",
        content=CSV_CONTENT,
        content_type="text/csv",
        conversation_id="docker-session",
    )
    reloaded_store = AttachmentStore(base_dir=temp_store.base_dir)
    reloaded = reloaded_store.get(metadata.attachment_id)
    assert reloaded is not None
    assert reloaded.original_filename == "pickup_report.csv"
    assert Path(reloaded.storage_path).exists()


def test_attachment_visible_across_service_instances(temp_store: AttachmentStore) -> None:
    """Regression: uploads from one service instance must be readable by another."""
    uploader = AttachmentService(store=temp_store)
    metadata = uploader.upload(
        filename="pickup_report.xlsx",
        content=_make_excel_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        conversation_id="session-a",
    )
    reader = AttachmentService(store=AttachmentStore(base_dir=temp_store.base_dir))
    processed = reader.get_processed(metadata.attachment_id)
    assert processed.file_type == "excel"
    assert processed.statistics["row_count"] == 2


def test_plan_data_sources_for_attachment_and_sql() -> None:
    sources = plan_data_sources(
        "Compare this report with today's database",
        has_attachments=True,
        file_types=["csv"],
    )
    assert DataSource.ATTACHMENT_AND_SQL in sources


def test_plan_data_sources_for_attachment_and_rag() -> None:
    sources = plan_data_sources(
        "Does this follow our SOP?",
        has_attachments=True,
        file_types=["pdf"],
    )
    assert DataSource.ATTACHMENT_AND_RAG in sources


def test_plan_data_sources_for_image_database_comparison() -> None:
    sources = plan_data_sources(
        "Compare it with today's database",
        has_attachments=True,
        file_types=["image"],
    )
    assert DataSource.ATTACHMENT_AND_SQL in sources


def test_attachment_memory_resolves_this_file_reference() -> None:
    memory = AttachmentMemory()
    metadata = AttachmentMetadata(
        filename="safe.csv",
        original_filename="pickup_report.csv",
        file_type="csv",
        file_size=10,
        storage_path="/tmp/safe.csv",
    )
    processed = process_csv(metadata, CSV_CONTENT)
    memory.register_contexts([processed])
    contexts, file_context = memory.resolve_for_question("Summarize this file")
    assert contexts[0].filename == "pickup_report.csv"
    assert "pickup_report.csv" in file_context["filenames"]


def test_file_reference_resolution_for_csv() -> None:
    left = process_csv(_metadata("monday.csv", "csv"), MONDAY_CSV)
    right = process_csv(_metadata("tuesday.csv", "csv"), TUESDAY_CSV)
    resolved = resolve_file_reference("Compare the CSV", [left, right])
    assert resolved[0].filename == "monday.csv"


def test_file_context_builder_includes_statistics() -> None:
    processed = process_csv(_metadata("pickup_report.csv", "csv"), CSV_CONTENT)
    context = build_file_context(question="Analyze this report", contexts=[processed])
    assert context["filenames"] == ["pickup_report.csv"]
    assert context["statistics"][0]["row_count"] == 3


@patch("tools.files.analyzer.execute")
@patch("tools.files.analyzer.get_latest_business_date")
def test_attachment_sql_comparison_uses_database_rows(
    mock_latest_date: MagicMock,
    mock_execute: MagicMock,
) -> None:
    from datetime import date

    mock_latest_date.return_value = date(2026, 6, 29)
    mock_execute.return_value = [
        {"hub": "Chicago Hub", "pickup_count": 12},
        {"hub": "New York Hub", "pickup_count": 98},
    ]
    processed = process_csv(_metadata("pickup_report.csv", "csv"), CSV_CONTENT)
    response = file_analyzer.analyze_attachments(
        question="Compare this report with today's database",
        resolved_question="Compare this report with today's database",
        contexts=[processed],
        file_context=build_file_context(question="Compare", contexts=[processed]),
        data_sources=[DataSource.ATTACHMENT_AND_SQL],
        intent="FILE_DATABASE_COMPARISON",
    )
    assert response is not None
    assert "GOFO operational database" in response.answer
    mock_execute.assert_called_once()


@patch("tools.files.analyzer.rag_answer")
def test_attachment_rag_comparison_uses_sop_retrieval(mock_rag: MagicMock) -> None:
    mock_rag.return_value = MagicMock(
        answer="Drivers must report failed pickups within 2 hours.",
        sources=[],
    )
    processed = process_docx(_metadata("process.docx", "docx"), _make_docx_bytes())
    response = file_analyzer.analyze_attachments(
        question="Does this follow our SOP?",
        resolved_question="Does this follow our SOP?",
        contexts=[processed],
        file_context=build_file_context(question="SOP", contexts=[processed]),
        data_sources=[DataSource.ATTACHMENT_AND_RAG],
        intent="FILE_RAG_COMPARISON",
    )
    assert response is not None
    assert "SOP" in response.answer
    mock_rag.assert_called_once()


def test_multi_file_comparison_detects_hub_changes() -> None:
    left = process_csv(_metadata("monday.csv", "csv"), MONDAY_CSV)
    right = process_csv(_metadata("tuesday.csv", "csv"), TUESDAY_CSV)
    response = file_analyzer.analyze_attachments(
        question="Compare these two reports",
        resolved_question="Compare these two reports",
        contexts=[left, right],
        file_context=build_file_context(question="Compare", contexts=[left, right]),
        data_sources=[DataSource.ATTACHMENT],
        intent="FILE_COMPARISON",
    )
    assert response is not None
    assert "monday.csv" in response.answer
    assert "tuesday.csv" in response.answer


@patch("tools.memory.conversation.extract_entities")
@patch("core.route_dispatcher.ask_core")
def test_csv_conversation_flow_uses_attachment_memory(
    mock_ask_core: MagicMock,
    mock_extract_entities: MagicMock,
    temp_store: AttachmentStore,
) -> None:
    mock_extract_entities.return_value = {}
    mock_ask_core.side_effect = _unexpected_core_call
    service = AttachmentService(store=temp_store)
    agent = GOFOAgent()
    agent.attachment_service = service

    metadata = service.upload(
        filename="pickup_report.csv",
        content=CSV_CONTENT,
        content_type="text/csv",
        conversation_id="session-csv",
    )
    first = agent.ask("Analyze this report", attachment_ids=[metadata.attachment_id])
    second = agent.ask("Which hub is worst?")
    third = agent.ask("Why?")
    fourth = agent.ask("What should operations do?")

    assert "pickup_report.csv" in first["answer"]
    assert "Chicago Hub" in second["answer"]
    assert "Root Cause" in third["answer"] or "Chicago Hub" in third["answer"]
    assert "Recommendations" in fourth["answer"] or "Recommendation" in fourth["answer"]
    assert mock_ask_core.call_count == 0


@patch("tools.memory.conversation.extract_entities")
@patch("core.route_dispatcher.ask_core")
def test_excel_conversation_flow_preserves_driver_context(
    mock_ask_core: MagicMock,
    mock_extract_entities: MagicMock,
    temp_store: AttachmentStore,
) -> None:
    mock_extract_entities.return_value = {}
    mock_ask_core.side_effect = _unexpected_core_call
    service = AttachmentService(store=temp_store)
    agent = GOFOAgent()
    agent.attachment_service = service

    metadata = service.upload(
        filename="pickup_report.xlsx",
        content=_make_excel_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        conversation_id="session-xlsx",
    )
    first = agent.ask(
        "Which driver has the highest performance?",
        attachment_ids=[metadata.attachment_id],
    )
    second = agent.ask("Which hub does he belong to?")
    third = agent.ask("Show his records")

    assert "Drew Nguyen" in first["answer"]
    assert "ORD Hub" in second["answer"] or "hub" in second["answer"].lower()
    assert "Drew Nguyen" in third["answer"] or "records" in third["answer"].lower()
    assert mock_ask_core.call_count == 0


@patch("tools.files.image_processor._analyze_image")
@patch("tools.memory.conversation.extract_entities")
@patch("core.route_dispatcher.ask_core")
@patch("tools.files.analyzer.execute")
@patch("tools.files.analyzer.get_latest_business_date")
def test_image_then_database_comparison_flow(
    mock_latest_date: MagicMock,
    mock_execute: MagicMock,
    mock_ask_core: MagicMock,
    mock_extract_entities: MagicMock,
    mock_analyze_image: MagicMock,
    temp_store: AttachmentStore,
) -> None:
    from datetime import date

    mock_extract_entities.return_value = {}
    mock_ask_core.side_effect = _unexpected_core_call
    mock_analyze_image.return_value = {
        "description": "Screenshot shows elevated delays at Chicago Hub.",
        "analysis": "Delay issue visible.",
    }
    mock_latest_date.return_value = date(2026, 6, 29)
    mock_execute.return_value = [{"hub": "Chicago Hub", "pickup_count": 10}]
    service = AttachmentService(store=temp_store)
    agent = GOFOAgent()
    agent.attachment_service = service

    metadata = service.upload(
        filename="dashboard.png",
        content=_make_png_bytes(),
        content_type="image/png",
        conversation_id="session-image",
    )
    first = agent.ask(
        "What operational issue do you see?",
        attachment_ids=[metadata.attachment_id],
    )
    second = agent.ask("Compare it with today's database")

    assert "dashboard.png" in first["answer"]
    assert "GOFO operational database" in second["answer"]
    assert mock_ask_core.call_count == 0


@patch("tools.files.analyzer.rag_answer")
@patch("tools.memory.conversation.extract_entities")
@patch("core.route_dispatcher.ask_core")
def test_pdf_conversation_flow_with_rag_comparison(
    mock_ask_core: MagicMock,
    mock_extract_entities: MagicMock,
    mock_rag: MagicMock,
    temp_store: AttachmentStore,
) -> None:
    mock_extract_entities.return_value = {}
    mock_ask_core.side_effect = _unexpected_core_call
    mock_rag.return_value = MagicMock(
        answer="Pickup failures must be escalated to dispatch.",
        sources=[],
    )
    service = AttachmentService(store=temp_store)
    agent = GOFOAgent()
    agent.attachment_service = service

    metadata = service.upload(
        filename="operations_report.pdf",
        content=_make_pdf_bytes(),
        content_type="application/pdf",
        conversation_id="session-pdf",
    )
    first = agent.ask("Summarize this", attachment_ids=[metadata.attachment_id])
    second = agent.ask("What are the biggest risks?")
    third = agent.ask("Does this follow our SOP?")

    assert "operations_report.pdf" in first["answer"]
    assert "Risk" in second["answer"] or "risk" in second["answer"].lower()
    assert "SOP" in third["answer"]
    assert mock_ask_core.call_count == 0


@patch("tools.memory.conversation.extract_entities")
@patch("core.route_dispatcher.ask_core")
def test_multi_csv_conversation_comparison_flow(
    mock_ask_core: MagicMock,
    mock_extract_entities: MagicMock,
    temp_store: AttachmentStore,
) -> None:
    mock_extract_entities.return_value = {}
    mock_ask_core.side_effect = _unexpected_core_call
    service = AttachmentService(store=temp_store)
    agent = GOFOAgent()
    agent.attachment_service = service

    monday = service.upload(
        filename="monday.csv",
        content=MONDAY_CSV,
        content_type="text/csv",
        conversation_id="session-multi",
    )
    tuesday = service.upload(
        filename="tuesday.csv",
        content=TUESDAY_CSV,
        content_type="text/csv",
        conversation_id="session-multi",
    )
    first = agent.ask(
        "Compare these reports",
        attachment_ids=[monday.attachment_id, tuesday.attachment_id],
    )
    second = agent.ask("Which hub changed the most?")
    third = agent.ask("Why?")

    assert "monday.csv" in first["answer"]
    assert "tuesday.csv" in first["answer"]
    assert "changed the most" in second["answer"]
    assert "Chicago Hub" in third["answer"] or "New York Hub" in third["answer"] or "Root Cause" in third["answer"]
    assert mock_ask_core.call_count == 0


def test_handle_error_returns_file_message_for_file_error() -> None:
    message = handle_error(FileError("The requested attachment could not be found."))
    assert message == "The requested attachment could not be found."


@patch("tools.memory.conversation.extract_entities")
@patch("core.route_dispatcher.ask_core")
def test_api_upload_then_ask_finds_attachment(
    mock_ask_core: MagicMock,
    mock_extract_entities: MagicMock,
    temp_store: AttachmentStore,
    monkeypatch,
) -> None:
    """Regression: API upload and session ask must share attachment metadata."""
    mock_extract_entities.return_value = {}
    mock_ask_core.side_effect = _unexpected_core_call
    shared_service = AttachmentService(store=temp_store)
    monkeypatch.setattr("api.server.attachment_service", shared_service)

    def _factory() -> GOFOAgent:
        session_agent = GOFOAgent()
        session_agent.attachment_service = shared_service
        return session_agent

    monkeypatch.setattr("api.server._create_session_agent", _factory)
    mock_session_manager = MagicMock()
    mock_session_manager.get_session.side_effect = lambda session_id: _factory()
    monkeypatch.setattr("api.server.session_manager", mock_session_manager)

    client = TestClient(app)
    upload_response = client.post(
        "/attachments",
        data={"session_id": "upload-session"},
        files={
            "files": (
                "order_report.xlsx",
                _make_excel_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload_response.status_code == 200
    attachment_id = upload_response.json()["attachments"][0]["attachment_id"]

    ask_response = client.post(
        "/ask",
        json={
            "question": "analyze this file, tell me what you discovered?",
            "attachments": [attachment_id],
            "session_id": "upload-session",
        },
    )
    assert ask_response.status_code == 200
    assert "order_report.xlsx" in ask_response.json()["answer"]
    assert mock_ask_core.call_count == 0


def test_post_attachments_endpoint(temp_store: AttachmentStore, monkeypatch) -> None:
    monkeypatch.setattr("api.server.attachment_service", AttachmentService(store=temp_store))
    client = TestClient(app)
    response = client.post(
        "/attachments",
        data={"session_id": "upload-session"},
        files={"files": ("pickup_report.csv", CSV_CONTENT, "text/csv")},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["attachments"][0]["original_filename"] == "pickup_report.csv"
    assert payload["attachments"][0]["processing_status"] == "UPLOADED"


def test_post_attachments_rejects_invalid_type(temp_store: AttachmentStore, monkeypatch) -> None:
    monkeypatch.setattr("api.server.attachment_service", AttachmentService(store=temp_store))
    client = TestClient(app)
    response = client.post(
        "/attachments",
        data={"session_id": "upload-session"},
        files={"files": ("malware.exe", b"MZ", "application/octet-stream")},
    )
    assert response.status_code == 400


def test_post_ask_backwards_compatible_without_attachments() -> None:
    client = TestClient(app)
    with patch("api.server.session_manager") as mock_session_manager:
        mock_session_agent = MagicMock()
        mock_session_manager.get_session.return_value = mock_session_agent
        mock_session_agent.ask.return_value = {
            "answer": "GOFO operations are stable.",
            "sources": [],
            "sql": "",
            "data": [],
            "analysis": {},
            "recommendations": [],
            "kpi": {},
            "raw": {},
        }
        response = client.post("/ask", json={"question": "How are operations today?"})
    assert response.status_code == 200
    mock_session_agent.ask.assert_called_once_with("How are operations today?", attachment_ids=[])


def test_post_ask_accepts_attachment_ids() -> None:
    client = TestClient(app)
    with patch("api.server.session_manager") as mock_session_manager:
        mock_session_agent = MagicMock()
        mock_session_manager.get_session.return_value = mock_session_agent
        mock_session_agent.ask.return_value = {
            "answer": "Analyzed report.",
            "sources": [],
            "sql": "",
            "data": [],
            "analysis": {"attachment_ids": ["abc"]},
            "recommendations": [],
            "kpi": {},
            "raw": {},
        }
        response = client.post(
            "/ask",
            json={"question": "Analyze this report", "attachments": ["abc"], "session_id": "s1"},
        )
    assert response.status_code == 200
    mock_session_agent.ask.assert_called_once_with("Analyze this report", attachment_ids=["abc"])


def _unexpected_core_call(*_args, **_kwargs) -> QueryResponse:
    raise AssertionError("ask_core should not be called for attachment-only analysis")

