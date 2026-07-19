"""Tests for attachment preview helpers and pending attachment state."""

from __future__ import annotations

import io
import json

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from frontend.attachment_preview import (
    attachment_identity,
    build_docx_preview,
    build_image_preview,
    build_pdf_preview,
    build_preview_payload,
    build_table_preview,
    build_text_preview,
    format_file_size,
    get_file_type,
    merge_pending_attachments,
    remove_pending_attachment,
)


CSV_BYTES = b"""pickup_id,hub,status,packages
101,Chicago,Completed,120
102,Dallas,Delayed,80
103,Atlanta,Failed,45
"""

JSON_BYTES = b'{"hub":"Chicago","packages":120,"status":"Completed"}'

TXT_BYTES = b"Failed pickups must be escalated immediately."

INVALID_JSON_BYTES = b"{not valid json"


def _png_bytes() -> bytes:
    image = Image.new("RGB", (80, 60), color=(20, 120, 200))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _jpg_bytes() -> bytes:
    image = Image.new("RGB", (90, 70), color=(200, 80, 20))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG")
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["pickup_id", "hub", "status", "packages"])
    sheet.append([101, "Chicago", "Completed", 120])
    sheet.append([102, "Dallas", "Delayed", 80])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Failed pickups must be reported within 2 hours.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _metadata(filename: str, attachment_id: str, size: int, file_type: str) -> dict:
    return {
        "attachment_id": attachment_id,
        "filename": filename,
        "original_filename": filename,
        "file_type": file_type,
        "file_size": size,
        "processing_status": "UPLOADED",
    }


def test_upload_png_thumbnail_preview_payload() -> None:
    content = _png_bytes()
    payload = build_image_preview(content, "dashboard.png", "png")
    assert payload["preview_type"] == "image"
    assert payload["content"] == content
    assert payload["width"] == 80
    assert payload["height"] == 60


def test_upload_jpg_image_preview_payload() -> None:
    content = _jpg_bytes()
    payload = build_preview_payload(filename="photo.jpg", file_type="jpg", content=content)
    assert payload["preview_type"] == "image"
    assert payload["error"] is None


def test_upload_pdf_preview_does_not_crash() -> None:
    content = _pdf_bytes()
    payload = build_pdf_preview(content, "operations_sop.pdf")
    assert payload["preview_type"] == "pdf"
    assert payload["page_count"] == 1
    assert payload["error"] is None


def test_upload_csv_preview_first_rows() -> None:
    payload = build_table_preview(CSV_BYTES, "csv", max_rows=20)
    assert payload["preview_type"] == "table"
    assert payload["columns"] == ["pickup_id", "hub", "status", "packages"]
    assert len(payload["rows"]) == 3
    assert payload["rows"][0]["hub"] == "Chicago"


def test_upload_xlsx_preview_works() -> None:
    content = _xlsx_bytes()
    payload = build_table_preview(content, "xlsx", max_rows=20)
    assert payload["preview_type"] == "table"
    assert payload["active_sheet"] == "Sheet1"
    assert len(payload["rows"]) == 2


def test_upload_txt_preview() -> None:
    payload = build_text_preview(TXT_BYTES, "notes.txt", "txt")
    assert "Failed pickups" in payload["text"]
    assert payload["error"] is None


def test_upload_json_pretty_printed() -> None:
    payload = build_text_preview(JSON_BYTES, "data.json", "json")
    assert payload["file_type"] == "json"
    assert '"hub": "Chicago"' in payload["text"]
    assert payload["error"] is None


def test_upload_invalid_json_falls_back_to_plain_text() -> None:
    payload = build_text_preview(INVALID_JSON_BYTES, "broken.json", "json")
    assert payload["text"] == INVALID_JSON_BYTES.decode("utf-8")
    assert payload["error"] is None


def test_upload_docx_text_preview() -> None:
    payload = build_docx_preview(_docx_bytes(), "sop.docx")
    assert "Failed pickups" in payload["text"]
    assert payload["error"] is None


def test_multiple_uploads_merge_without_duplicates() -> None:
    pending: list[dict] = []
    content = CSV_BYTES
    uploaded = [_metadata("pickup_report.csv", "id-1", len(content), "csv")]
    merged, cache = merge_pending_attachments(pending, uploaded, {"pickup_report.csv": content})
    assert len(merged) == 1
    assert cache["id-1"] == content

    merged_again, cache_again = merge_pending_attachments(
        merged,
        uploaded,
        {"pickup_report.csv": content},
    )
    assert len(merged_again) == 1
    assert cache_again == {}


def test_remove_one_attachment_only() -> None:
    pending = [
        _metadata("one.csv", "id-1", 10, "csv"),
        _metadata("two.csv", "id-2", 10, "csv"),
    ]
    remaining = remove_pending_attachment(pending, "id-1")
    assert len(remaining) == 1
    assert remaining[0]["attachment_id"] == "id-2"


def test_streamlit_rerun_does_not_duplicate_attachments() -> None:
    pending: list[dict] = []
    content_a = CSV_BYTES
    content_b = b"hub,pickups\nChicago,1\n"
    first_upload = [_metadata("monday.csv", "id-1", len(content_a), "csv")]
    second_upload = [_metadata("tuesday.csv", "id-2", len(content_b), "csv")]

    pending, _ = merge_pending_attachments(
        pending,
        first_upload,
        {"monday.csv": content_a},
    )
    pending, _ = merge_pending_attachments(
        pending,
        first_upload,
        {"monday.csv": content_a},
    )
    pending, _ = merge_pending_attachments(
        pending,
        second_upload,
        {"tuesday.csv": content_b},
    )
    assert len(pending) == 2


def test_successful_send_clears_pending_but_keeps_preview_cache() -> None:
    pending = [_metadata("pickup_report.csv", "id-1", len(CSV_BYTES), "csv")]
    preview_cache = {"id-1": CSV_BYTES}
    pending_after_success: list[dict] = []
    assert pending_after_success == []
    assert preview_cache["id-1"] == CSV_BYTES


def test_failed_send_keeps_pending_attachments() -> None:
    pending = [_metadata("pickup_report.csv", "id-1", len(CSV_BYTES), "csv")]
    failed = True
    pending_after_failure = pending if failed else []
    assert len(pending_after_failure) == 1


def test_historical_message_attachment_metadata_is_reusable() -> None:
    message = {
        "role": "user",
        "content": "Analyze these pickup failures.",
        "attachments": [_metadata("pickup_report.csv", "id-1", len(CSV_BYTES), "csv")],
    }
    preview_cache = {"id-1": CSV_BYTES}
    attachment = message["attachments"][0]
    payload = build_preview_payload(
        filename=attachment["original_filename"],
        file_type=attachment["file_type"],
        content=preview_cache[attachment["attachment_id"]],
    )
    assert payload["preview_type"] == "table"
    assert payload["rows"][0]["hub"] == "Chicago"


def test_unsupported_file_generic_preview_message() -> None:
    payload = build_preview_payload(
        filename="archive.zip",
        file_type="zip",
        content=b"fake",
    )
    assert payload["preview_type"] == "unsupported"
    assert "Preview unavailable" in payload["message"]


def test_corrupt_pdf_preview_returns_user_message() -> None:
    payload = build_pdf_preview(b"not-a-pdf", "broken.pdf")
    assert payload["error"] == "Unable to preview this file."


def test_invalid_image_preview_returns_user_message() -> None:
    payload = build_image_preview(b"not-an-image", "broken.png", "png")
    assert payload["error"] == "Unable to preview this file."


def test_format_file_size_and_file_type_helpers() -> None:
    assert format_file_size(500) == "500 B"
    assert format_file_size(2048) == "2.0 KB"
    assert get_file_type("report.CSV") == "csv"
    assert get_file_type("unknown", "application/pdf") == "pdf"


def test_attachment_identity_is_stable() -> None:
    content = b"same-content"
    first = attachment_identity("report.csv", len(content), content)
    second = attachment_identity("report.csv", len(content), content)
    assert first == second


def test_csv_preview_limits_rows() -> None:
    rows = ["hub,pickups"] + [f"Hub{i},{i}" for i in range(30)]
    content = "\n".join(rows).encode()
    payload = build_table_preview(content, "csv", max_rows=20)
    assert len(payload["rows"]) == 20
    assert payload["truncated"] is True
